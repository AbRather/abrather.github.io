from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

import yaml
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


REPO_ROOT = Path(__file__).resolve().parents[1]
OUTPUT_FILE = REPO_ROOT / "presentation" / "hamid_rather_introduction.pptx"

NAVY = RGBColor(10, 25, 48)
CYAN = RGBColor(52, 191, 255)
WHITE = RGBColor(255, 255, 255)
LIGHT_TEXT = RGBColor(219, 233, 245)


@dataclass
class SourceData:
    config: dict[str, Any]
    careers: list[dict[str, Any]]
    education: list[dict[str, Any]]
    projects: list[dict[str, Any]]
    skills: dict[str, Any]
    contact: dict[str, Any]


def load_yaml(path: Path) -> Any:
    with path.open("r", encoding="utf-8") as fh:
        return yaml.safe_load(fh)


def load_data() -> SourceData:
    return SourceData(
        config=load_yaml(REPO_ROOT / "_config.yml"),
        careers=load_yaml(REPO_ROOT / "_data" / "careers.yml"),
        education=load_yaml(REPO_ROOT / "_data" / "education.yml"),
        projects=load_yaml(REPO_ROOT / "_data" / "projects.yml"),
        skills=load_yaml(REPO_ROOT / "_data" / "skills.yml"),
        contact=load_yaml(REPO_ROOT / "_data" / "contact.yml"),
    )


def style_slide_background(slide) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = NAVY


def add_title(
    slide,
    title: str,
    subtitle: str | None = None,
    website: str | None = None,
) -> None:
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.6), Inches(11.8), Inches(1.2))
    title_tf = title_box.text_frame
    title_tf.clear()
    p = title_tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = WHITE

    if subtitle:
        subtitle_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(11.8), Inches(0.8))
        sub_tf = subtitle_box.text_frame
        sub_tf.clear()
        p = sub_tf.paragraphs[0]
        p.text = subtitle
        p.font.size = Pt(23)
        p.font.color.rgb = CYAN

    if website:
        website_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.45), Inches(11.8), Inches(0.5))
        w_tf = website_box.text_frame
        w_tf.clear()
        p = w_tf.paragraphs[0]
        p.text = website
        p.font.size = Pt(16)
        p.font.color.rgb = LIGHT_TEXT

    line = slide.shapes.add_shape(1, Inches(0.8), Inches(3.0), Inches(4.0), Inches(0.05))
    line.fill.solid()
    line.fill.fore_color.rgb = CYAN
    line.line.fill.background()


def add_bullets(slide, items: list[str], left: float = 0.9, top: float = 2.0, width: float = 11.4, height: float = 4.9) -> None:
    box = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True

    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.level = 0
        p.font.size = Pt(21)
        p.font.color.rgb = LIGHT_TEXT
        p.space_after = Pt(10)


def add_two_column_sections(slide, left_title: str, left_items: list[str], right_title: str, right_items: list[str]) -> None:
    left_head = slide.shapes.add_textbox(Inches(0.8), Inches(1.9), Inches(5.8), Inches(0.5))
    p = left_head.text_frame.paragraphs[0]
    p.text = left_title
    p.font.bold = True
    p.font.size = Pt(20)
    p.font.color.rgb = CYAN

    right_head = slide.shapes.add_textbox(Inches(7.0), Inches(1.9), Inches(5.5), Inches(0.5))
    p = right_head.text_frame.paragraphs[0]
    p.text = right_title
    p.font.bold = True
    p.font.size = Pt(20)
    p.font.color.rgb = CYAN

    add_bullets(slide, left_items, left=0.8, top=2.4, width=5.8, height=4.6)
    add_bullets(slide, right_items, left=7.0, top=2.4, width=5.5, height=4.6)


def find_career(careers: list[dict[str, Any]], needle: str) -> dict[str, Any]:
    needle_low = needle.lower()
    for role in careers:
        detail = role.get("name", {}).get("detail", "")
        if needle_low in detail.lower():
            return role
    raise ValueError(f"Could not find career entry containing: {needle}")


def find_project(projects: list[dict[str, Any]], needle: str) -> dict[str, Any]:
    needle_low = needle.lower()
    for project in projects:
        if needle_low in project.get("name", "").lower():
            return project
    raise ValueError(f"Could not find project entry containing: {needle}")


def add_notes(slide, notes: str) -> None:
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = notes


def create_presentation() -> Path:
    data = load_data()
    author = data.config["author"]

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    # 1. Title
    slide = prs.slides.add_slide(blank)
    style_slide_background(slide)
    add_title(
        slide,
        "Hello, I’m Hamid Rather",
        "Aerospace Engineer | Data Analyst | Researcher",
        "https://abrather.github.io/",
    )
    add_notes(slide, "Short intro and purpose: give coworkers a quick overview of background, projects, and how I can contribute.")

    # 2. Professional profile
    slide = prs.slides.add_slide(blank)
    style_slide_background(slide)
    add_title(slide, "Professional Profile")
    add_bullets(
        slide,
        [
            "Aerospace and mechanical engineering background.",
            "Works at the hardware-software intersection.",
            "Focus areas: composite materials, cryogenics, and data analysis.",
            "Designs experiments, integrates equipment, and automates analysis workflows.",
        ],
    )

    # 3. Current role
    current = find_career(data.careers, "Faserinstitut Bremen")
    current_date = current.get("date", {}).get("detail", "")
    current_details = current.get("desc", {}).get("details", [])

    slide = prs.slides.add_slide(blank)
    style_slide_background(slide)
    add_title(slide, "Current Role", f"Student Researcher / Werkstudent — Faserinstitut Bremen ({current_date})")
    add_bullets(
        slide,
        [
            current_details[0],
            current_details[1],
            current_details[2],
            current_details[3],
            current_details[4],
        ],
    )
    add_notes(slide, "Emphasize practical lab ownership plus simulation and Python-based data infrastructure.")

    # 4. Research focus
    hytank = find_project(data.projects, "hytank")
    slide = prs.slides.add_slide(blank)
    style_slide_background(slide)
    add_title(slide, "Research Focus", "Composite materials for cryogenic hydrogen storage")
    add_two_column_sections(
        slide,
        "Why this matters",
        [
            "Hydrogen tanks must be lightweight and structurally reliable.",
            "Cryogenic conditions challenge material durability and sealing behavior.",
            "Permeability and damage behavior directly affect performance and safety.",
        ],
        "Current research contribution",
        [
            "Cryogenic material characterization and permeation testing.",
            "Design/integration of high-pressure helium test equipment.",
            "Coupled simulation and Python-based analysis workflows.",
        ],
    )
    add_notes(slide, f"Reference project: {hytank.get('name')}.")

    # 5. Current projects
    wagupa = find_project(data.projects, "wagupa")
    benene = find_project(data.projects, "benene")
    slide = prs.slides.add_slide(blank)
    style_slide_background(slide)
    add_title(slide, "Current Projects")
    add_two_column_sections(
        slide,
        hytank["name"],
        [
            hytank.get("desc", ""),
            *hytank.get("features", [])[:2],
        ],
        wagupa["name"],
        [
            wagupa.get("desc", ""),
            *wagupa.get("features", [])[:2],
        ],
    )

    benene_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.0), Inches(12.0), Inches(1.1))
    benene_tf = benene_box.text_frame
    benene_tf.clear()
    p = benene_tf.paragraphs[0]
    p.text = f"{benene['name']}: {benene.get('desc', '')}"
    p.font.size = Pt(16)
    p.font.color.rgb = LIGHT_TEXT

    # 6. Previous experience
    polaris = find_career(data.careers, "Polaris")
    rfa = find_career(data.careers, "Rocket Factory Augsburg")
    slide = prs.slides.add_slide(blank)
    style_slide_background(slide)
    add_title(slide, "Previous Experience")
    add_two_column_sections(
        slide,
        f"{polaris['name']['detail']} ({polaris['date']['detail']})",
        polaris["desc"]["details"][:4],
        f"{rfa['name']['detail']} ({rfa['date']['detail']})",
        rfa["desc"]["details"][:4],
    )

    # 7. Earlier experience
    spantrik = find_career(data.careers, "SpanTrik")
    srm = find_career(data.careers, "Advanced Computational")
    slide = prs.slides.add_slide(blank)
    style_slide_background(slide)
    add_title(slide, "Earlier Experience")
    add_two_column_sections(
        slide,
        f"{spantrik['name']['detail']} ({spantrik['date']['detail']})",
        spantrik["desc"]["details"],
        f"{srm['name']['detail']} ({srm['date']['detail']})",
        srm["desc"]["details"][:4],
    )

    # 8. Technical skills
    skill_data = data.skills
    slide = prs.slides.add_slide(blank)
    style_slide_background(slide)
    add_title(slide, "Technical Skills")
    add_two_column_sections(
        slide,
        "Programming & Data",
        [
            ", ".join([s["name"] for s in skill_data.get("circular_skills", [])[:3]]),
            ", ".join(skill_data.get("technical_list", [])[:4]),
        ],
        "Engineering, Simulation & Tools",
        [
            ", ".join([s["name"] for s in skill_data.get("circular_skills", [])[3:6]]),
            ", ".join(skill_data.get("knowledge_areas", [])[:4]),
            ", ".join(skill_data.get("technical_list", [])[4:8]),
        ],
    )

    # 9. Engineering approach
    slide = prs.slides.add_slide(blank)
    style_slide_background(slide)
    add_title(slide, "Engineering Approach")

    steps = [
        ("1. Understand", "Clarify objective, constraints, and success criteria."),
        ("2. Design", "Develop hardware, simulation, and analysis workflows."),
        ("3. Build & Integrate", "Connect systems, sensors, and data capture."),
        ("4. Test & Analyze", "Run experiments and extract reliable insights."),
        ("5. Communicate", "Present outcomes with clear visuals and recommendations."),
    ]
    y = 2.0
    for head, text in steps:
        h = slide.shapes.add_textbox(Inches(1.0), Inches(y), Inches(3.0), Inches(0.45))
        hp = h.text_frame.paragraphs[0]
        hp.text = head
        hp.font.bold = True
        hp.font.size = Pt(18)
        hp.font.color.rgb = CYAN

        body = slide.shapes.add_textbox(Inches(3.4), Inches(y), Inches(8.9), Inches(0.45))
        bp = body.text_frame.paragraphs[0]
        bp.text = text
        bp.font.size = Pt(18)
        bp.font.color.rgb = LIGHT_TEXT
        y += 0.9

    # 10. Education and languages
    edu = [e for e in data.education if "Universität Bremen" in e.get("school", "") or "SRM University" in e.get("school", "")]
    langs = skill_data.get("language_skills", [])

    slide = prs.slides.add_slide(blank)
    style_slide_background(slide)
    add_title(slide, "Education and Languages")
    add_two_column_sections(
        slide,
        "Education",
        [
            f"{entry['degree']} — {entry['school']} ({entry['date']})"
            for entry in edu
        ],
        "Languages",
        [f"{lang['name']}: {lang['description']}" for lang in langs],
    )

    # 11. Personal interests
    sports = author.get("sports", "")
    hobbies = author.get("hobbies", "")
    slide = prs.slides.add_slide(blank)
    style_slide_background(slide)
    add_title(slide, "Personal Interests")
    add_bullets(
        slide,
        [
            "Outside work, I enjoy:",
            f"• {sports}",
            f"• {hobbies}",
        ],
        top=2.3,
    )

    # 12. Contribution goals
    slide = prs.slides.add_slide(blank)
    style_slide_background(slide)
    add_title(slide, "What I Hope to Contribute")
    add_two_column_sections(
        slide,
        "How I can support the team",
        [
            "Practical engineering and experimental execution.",
            "Data-driven problem solving with Python automation.",
            "Structured simulation and test-analysis workflows.",
        ],
        "What I am looking forward to",
        [
            "Collaborating across disciplines.",
            "Learning from ongoing projects and peers.",
            "Contributing to reliable, high-impact outcomes.",
        ],
    )

    # 13. Thank-you/contact
    slide = prs.slides.add_slide(blank)
    style_slide_background(slide)
    add_title(slide, "Thank You")

    contact_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(12.0), Inches(2.6))
    tf = contact_box.text_frame
    tf.clear()

    lines = [
        "Hamid Rather",
        "Aerospace Engineer | Data Analyst | Researcher",
        f"Website: {data.config['url']}/",
        f"Email: {data.contact['email']['value']}",
    ]

    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = PP_ALIGN.LEFT
        p.font.size = Pt(30 if i == 0 else 20)
        p.font.bold = i == 0
        p.font.color.rgb = WHITE if i == 0 else LIGHT_TEXT
        p.space_after = Pt(10)

    add_notes(slide, "Invite colleagues to reach out for collaboration on experiments, analysis, and simulation tasks.")

    OUTPUT_FILE.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUTPUT_FILE)
    return OUTPUT_FILE


if __name__ == "__main__":
    generated = create_presentation()
    print(f"Created {generated}")
